"""정규화된 주문 라인을 협력사별 엑셀 양식(templates/*.xlsx)에 채워 넣는다."""
import json
import re
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# 행 강조 색상: 수량 2개 이상 / 동일 수령인+주소 중복 / 둘 다 해당 시 각각 다른 색으로 구분
QTY_FILL = PatternFill(fill_type="solid", start_color="FFF9E48B", end_color="FFF9E48B")
DUP_FILL = PatternFill(fill_type="solid", start_color="FFBFE0F5", end_color="FFBFE0F5")
BOTH_FILL = PatternFill(fill_type="solid", start_color="FFF7C592", end_color="FFF7C592")
# 상품/상품명 참고표에 없어서 원본 그대로 채운 행(수기 확인 필요) 표시용 별색.
UNMATCHED_NAME_FILL = PatternFill(fill_type="solid", start_color="FFE0B3FF", end_color="FFE0B3FF")

BASE_DIR = Path(__file__).resolve().parent.parent

with open(BASE_DIR / "config" / "vendors.json", encoding="utf-8") as f:
    VENDORS = json.load(f)["vendors"]

with open(BASE_DIR / "config" / "column_aliases.json", encoding="utf-8") as f:
    _aliases_raw = json.load(f)

FIELD_TO_HEADERS = {k: v for k, v in _aliases_raw.items() if not k.startswith("_")}

_HEADER_TO_FIELD = {}
for _field, _names in FIELD_TO_HEADERS.items():
    for _name in _names:
        _HEADER_TO_FIELD[_name] = _field

_NO_HEADERS = {"NO", "No.", "No", "no", "번호", "순번"}


def _build_col_map(ws, cfg):
    """열 번호 -> ("static", 값) | ("combine",) | ("seq",) | ("year"/"month"/"day",)
    | ("shipment_flag",) | ("field", 표준필드명) | ("code",) | ("name_pair", "a"|"b")
    | ("consolidated", 표시문구) | ("group_seq",) 로 매핑한다. 매핑되지 않은 열은
    자동으로 채우지 않고 비워둔다(택배사/송장번호 등 나중에 수기로 채우는 열).
    (col_map, dup_cols) 튜플을 반환하며, dup_cols는 템플릿에 헤더가 실수로
    중복된 탓에 값을 채우지 못한 열 번호 목록이다(호출부가 이 열들을 숨김
    처리한다)."""
    header_row = cfg["header_row"]
    field_overrides = cfg.get("field_overrides", {})
    static_fields = cfg.get("static_fields", {})
    combine_field = cfg.get("combine_product_option_field")
    code_column = cfg.get("code_column")
    name_pair_columns = cfg.get("name_pair_columns", {})
    consolidated_col = cfg.get("consolidated_shipping_column")
    group_seq_col = cfg.get("group_seq_column")

    col_map = {}
    dup_cols = []
    seen = set()
    for c in range(1, ws.max_column + 1):
        header = ws.cell(row=header_row, column=c).value
        if not isinstance(header, str):
            continue
        header = header.strip()
        if not header:
            continue

        if header in static_fields:
            spec = ("static", static_fields[header])
        elif combine_field and header == combine_field:
            spec = ("combine",)
        elif code_column and header == code_column:
            spec = ("code",)
        elif consolidated_col and header == consolidated_col:
            spec = ("consolidated", consolidated_col)
        elif group_seq_col and header == group_seq_col:
            spec = ("group_seq",)
        elif header in name_pair_columns:
            spec = ("name_pair", name_pair_columns[header])
        elif header in field_overrides:
            spec = ("field", field_overrides[header])
        elif header in _NO_HEADERS:
            spec = ("seq",)
        elif header == "년":
            spec = ("year",)
        elif header == "월":
            spec = ("month",)
        elif header == "일":
            spec = ("day",)
        elif header == "배송건수":
            spec = ("shipment_flag",)
        elif header in _HEADER_TO_FIELD:
            spec = ("field", _HEADER_TO_FIELD[header])
        else:
            continue

        # 템플릿에 헤더가 실수로 중복된 열(예: "옵션"이 두 번)이 있으면
        # 첫 번째 열에만 값을 채우고 나머지는 숨김 처리한다(수식이 열
        # 위치를 참조하고 있어 템플릿에서 열 자체를 지우진 않음) — 같은
        # 값이 여러 열에 중복으로 찍히는 걸 막는다.
        if spec in seen:
            dup_cols.append(c)
            continue
        seen.add(spec)
        col_map[c] = spec
    return col_map, dup_cols


# ---------------------------------------------------------------------------
# 상품코드 조회: 협력사가 준 상품명/옵션명 <-> 자체 코드 매핑표(data/reference/
# codes_{협력사}.json)에서, 정규화한 이름이 주문 상품명+옵션 텍스트에 부분
# 일치하면 그 코드를 채운다. 매핑표 자체가 다른 상품군을 담고 있거나 커버리지가
# 낮을 수 있어 매치가 없으면 그냥 비워둔다(수기로 채우게).
# ---------------------------------------------------------------------------
_CODE_MATCH_MIN_LEN = 4
_code_tables = {}


def _normalize_for_code(s):
    s = str(s or "")
    s = re.sub(r'^[\(\[].*?[\)\]]', '', s)
    # 다른 매칭 로직(_normalize_for_match)과 같은 통일 규칙을 공유한다 —
    # "10개입"/"10입", "1박스"/"1세트" 같은 같은 뜻 다른 표기 차이 때문에
    # 상품코드 조회가 실패하지 않도록(cfg가 필요 없는 정규화만 적용).
    s = _normalize_container_words(s)
    s = _normalize_package_count(s)
    s = _normalize_stick_count(s)
    s = _strip_redundant_trailing_container(s)
    s = _normalize_candy_word(s)
    s = re.sub(r'[\s"\'/,.\-_★\[\]]+', '', s)
    return s.strip()


def _load_code_table(vendor_name, cfg):
    if vendor_name not in _code_tables:
        code_file = cfg.get("code_lookup_file")
        if not code_file:
            _code_tables[vendor_name] = None
        else:
            with open(BASE_DIR / code_file, encoding="utf-8") as f:
                entries = json.load(f)
            items = [(_normalize_for_code(e["name"]), e["code"]) for e in entries]
            items = [(n, c) for n, c in items if len(n) >= _CODE_MATCH_MIN_LEN]
            items.sort(key=lambda x: len(x[0]), reverse=True)
            _code_tables[vendor_name] = items
    return _code_tables[vendor_name]


def _code_for(vendor_name, cfg, rec):
    table = _load_code_table(vendor_name, cfg)
    if not table:
        return None
    text = _normalize_for_code(f"{rec.get('product_name') or ''} {rec.get('option') or ''}")
    if not text:
        return None
    for name, code in table:
        if name in text:
            return code
    return None


# ---------------------------------------------------------------------------
# 상품명/옵션명 쌍 조회: 협력사 자체 상품 목록(예: 이엑스 "상품"/"상품명" 두 컬럼
# 조합)을 data/reference/names_{협력사}.json({"a":.., "b":..} 쌍 목록)으로 갖고
# 있는 경우, 주문 상품명+옵션과 매칭되면 그 협력사 고유 표기(a, b)를 그대로 쓴다.
#
# 마켓 원본 상품명은 "카스텔크렘 포지타노 레몬 캔디 750g"처럼 협력사 내부
# 표기("레몬1 750g")와 단어 순서·구성이 달라 통짜 문자열 부분일치로는 못 찾는
# 경우가 많다. 대신 "a" 쪽에서 의미 있는 단어(한글 2글자 이상, 용량 표기 NNNg)를
# 전부 뽑아 그 단어들이 (순서 상관없이) 주문 텍스트에 전부 나타나는지로 판단한다
# — 예: "레몬1 750g" -> 필수 토큰 {"레몬","750g"} 둘 다 주문 텍스트에 있으면 매치.
# 후보가 여러 개면 필수 토큰이 더 많고(더 구체적인) 쪽을 우선한다.
# 매치가 없으면 우리 쪽 상품명을 그대로 양쪽 컬럼에 채운다(공란보다는 낫다).
# ---------------------------------------------------------------------------
_name_pair_tables = {}
# 'x'/'X'는 이 도메인에서 거의 항상 곱셈 기호("8개입x2박스")라, 단위 글자
# 범위에서 빼서 토큰이 그 앞에서 끊기게 한다(안 그러면 "8개입x2박스"가
# "8개입x"라는 뒤섞인 토큰이 돼 정식 옵션 목록의 "8입"과 매치가 안 된다).
_TOKEN_RE = re.compile(r'\d+[가-힣a-wyzA-WYZ]*|[가-힣]{2,}')


def _tokenize(s):
    """의미 있는 단어(한글 2글자 이상)와 단위 붙은 수량 표기(15포/750g/20티백
    등)를 뽑는다. 단위 없는 맨숫자(그냥 "1")는 어디에나 있을 수 있어 제외한다."""
    return {t for t in _TOKEN_RE.findall(s) if not t.isdigit()}


def _collect_tied_top(text, candidates, extra_ok=None):
    """candidates 중 필수 토큰이 전부 text에 있고(extra_ok가 있으면 그것도
    통과하는) 것들 중 가장 구체적인(토큰 글자수 합이 큰) 그룹을 전부 모아
    반환한다(동점 여러 개 가능). 하나도 없으면 빈 리스트."""
    best_score = None
    tied = []
    for tokens, payload in candidates:
        if not all(tok in text for tok in tokens):
            continue
        if extra_ok is not None and not extra_ok(tokens, payload):
            continue
        score = sum(len(t) for t in tokens)
        if best_score is None or score > best_score:
            best_score, tied = score, [(tokens, payload)]
        elif score == best_score:
            tied.append((tokens, payload))
    return tied


def _tiebreak(text, tied):
    """동점 후보가 여럿이면: 1) 원문에 "혼합"이 있으면 그 단어를 포함하는
    후보를 우선한다(예: "레몬오렌지 혼합캔디사탕"에서 "오렌지"가 우연히
    부분 문자열로 걸려도, 명시적으로 "혼합"이라고 쓰여 있으면 그게 더
    신뢰도 높은 신호다). 2) 그래도 여럿이면 후보 토큰 중 원문에서 가장
    먼저 나타나는 토큰의 위치가 앞선 후보를 고른다(예: "레몬/모로오렌지"
    처럼 대등하게 나열된 경우 먼저 언급된 쪽을 기본값으로 본다)."""
    if len(tied) == 1:
        return tied[0]
    if "혼합" in text:
        mixed = [c for c in tied if "혼합" in c[0]]
        if mixed:
            tied = mixed
            if len(tied) == 1:
                return tied[0]
    return min(tied, key=lambda c: min(text.find(t) for t in c[0]))


def _best_token_match_with_tokens(text, candidates):
    """candidates: [(tokens, payload), ...]. 필수 토큰이 전부 text에 있는 후보 중
    가장 구체적인(토큰 글자수 합이 큰) 것의 (tokens, payload)를 반환. 없으면 None.
    공백·기호 차이로 매칭이 실패하지 않도록("배&도라지" vs "배도라지") text에서
    공백과 흔한 연결기호를 제거하고 비교한다(토큰 쪽은 애초에 정규식이 그런
    문자를 건너뛰어 섞이지 않는다). 동점이면 _tiebreak로 하나를 고른다."""
    text = re.sub(r'[\s&+/,.\-_★]+', '', text)
    tied = _collect_tied_top(text, candidates)
    return _tiebreak(text, tied) if tied else None


def _best_token_match(text, candidates):
    result = _best_token_match_with_tokens(text, candidates)
    return result[1] if result else None


def _best_validated_match(text, candidates, identity_tokens):
    """_best_token_match_with_tokens와 같지만, 후보의 필수 토큰이 text에
    다 있다는 것뿐 아니라 identity_tokens(식별 단어 집합)가 그 후보에서도
    확인되는지 같이 본다. 후보 쪽은 토큰 집합이 아니라 후보의 원문(payload)
    자체를 공백·기호 다 지우고 이어붙인 문자열에 대해 부분 문자열로
    있는지 본다 — 토큰 단위로 비교하면 "배&도라지스틱"(원문, 기호+공백
    없음)과 "배도라지 스틱"(정식 옵션명, 공백 있음)처럼 똑같은 뜻인데
    토큰이 갈라지는 지점이 달라서("배도라지"+"스틱" vs "도라지스틱") 놓치는
    경우가 생긴다. 원문 이어붙인 문자열로 보면 "도라지스틱"이 "...배도라지
    스틱..."의 부분 문자열로 잡혀 정상 인정된다. 후보 원문이 문자열이
    아니면(예: 이엑스 "상품"/"상품명" 쌍) 두 값을 합쳐서 같은 방식으로 본다.
    동점이면 _tiebreak로 하나를 고른다 — 가장 구체적인 후보 하나만 보고
    포기하면, 상품명에 여러 맛 이름이 같이 있는 경우 엉뚱한 맛으로 매칭된
    걸 걸러내고도 실제로 맞는 다른 후보를 놓치게 된다."""
    text = re.sub(r'[\s&+/,.\-_★]+', '', text)

    def identity_ok(tokens, payload):
        payload_text = payload if isinstance(payload, str) else " ".join(payload)
        payload_flat = re.sub(r'[\s&+/,.\-_★]+', '', payload_text)
        payload_flat = _normalize_candy_word(payload_flat)
        return all(idt in payload_flat for idt in identity_tokens)

    tied = _collect_tied_top(text, candidates, extra_ok=identity_ok)
    return _tiebreak(text, tied) if tied else None


def _best_relaxed_match(text, candidates, identity_tokens):
    """_best_validated_match까지도 실패했을 때 마지막으로 시도하는 매칭.
    지금까지는 "후보의 모든 필수 토큰이 원문에 들어있어야" 인정했는데,
    반대로 고객이 실제로 적은 식별 단어(identity_tokens)가 후보 이름
    "안에" 포함돼 있는지만 본다 — 후보 쪽에 고객이 안 적은 브랜드/설명이
    더 붙어 있어도 상관없다(예: 원문 "모로오렌지"가 후보명 "모로오렌지
    캔디"의 부분 문자열이면 인정, 원문 "포켓몬"+"캔디머신"이 후보명 "포켓몬
    캔디머신 아이알파캔디"에 다 포함돼 있으면 원문에 없는 "아이알파캔디"는
    무시하고 인정). 사이즈(숫자+단위) 토큰만은 후보 쪽에도 정확히 있어야
    한다(200g짜리를 300g 후보로 착각하면 안 되므로). 이 조건을 만족하는
    후보가 정확히 하나뿐일 때만 인정한다 — 둘 이상이면 어느 쪽인지 구분이
    안 되는 것이므로 아무것도 고르지 않고 사람이 확인하게 한다."""
    if not identity_tokens:
        return None
    size_tokens = {t for t in _tokenize(text) if t[0].isdigit()}
    matched = []
    for tokens, payload in candidates:
        if not all(any(idt in ct for ct in tokens) for idt in identity_tokens):
            continue
        cand_size_tokens = {t for t in tokens if t[0].isdigit()}
        if size_tokens and cand_size_tokens and not (size_tokens & cand_size_tokens):
            continue
        matched.append((tokens, payload))
    if len(matched) == 1:
        return matched[0]
    return None


def _load_name_pair_table(vendor_name, cfg):
    if vendor_name not in _name_pair_tables:
        pair_file = cfg.get("name_pair_lookup_file")
        if not pair_file:
            _name_pair_tables[vendor_name] = None
        else:
            with open(BASE_DIR / pair_file, encoding="utf-8") as f:
                entries = json.load(f)
            items = []
            for e in entries:
                tokens = _tokenize(_normalize_for_match(e["a"], cfg))
                if tokens:
                    items.append((tokens, (e["a"], e["b"])))
            items.sort(key=lambda x: sum(len(t) for t in x[0]), reverse=True)
            _name_pair_tables[vendor_name] = items
    return _name_pair_tables[vendor_name]


def _name_pair_match(vendor_name, cfg, rec):
    """상품명/옵션 조합에 맞는 협력사 자체 표기("상품"/"상품명" 쌍)를 찾는다.
    옵션 텍스트 자체의 식별 단어가 후보에 없으면(예: 상품명이 "레몬/
    모로오렌지" 둘 다 소개하는 문구인데 고객은 실제로 "레몬"만 골랐을 때,
    상품명의 "모로오렌지"에 엉뚱하게 끌려 매칭되는 사고) 매칭을 인정하지
    않는다 — 정식 옵션 매칭(_find_option_match)과 같은 이유·같은 방식."""
    table = _load_name_pair_table(vendor_name, cfg)
    if not table:
        return None
    option_text = _normalize_for_match((rec.get("option") or "").strip(), cfg)
    combined_text = _normalize_for_match(
        f"{rec.get('product_name') or ''} {rec.get('option') or ''}", cfg
    )
    if not combined_text.strip():
        return None
    identity_tokens = {
        t for t in _tokenize(option_text)
        if not t[0].isdigit() and t not in _OPTION_BOILERPLATE_TOKENS
    }
    result = _best_validated_match(combined_text, table, identity_tokens)
    if not result:
        result = _best_relaxed_match(option_text, table, identity_tokens)
    if not result:
        return None
    _, payload = result
    return {"a": payload[0], "b": payload[1]}


# ---------------------------------------------------------------------------
# 옵션명 정규화: 협력사가 인정하는 옵션명 목록(예: "이플코리아 옵션명" 참고
# 시트, data/reference/options_{협력사}.json 문자열 리스트)이 있는 경우, 우리
# 쪽 옵션 텍스트가 두루뭉술해도 상품명+옵션 전체에서 그 목록의 항목과 토큰이
# 다 겹치면 그 정식 옵션명으로 바꿔 쓴다. 매치가 없으면 원래 옵션 텍스트를
# 그대로 둔다.
# ---------------------------------------------------------------------------
_option_canon_tables = {}
_REDUNDANT_ONE_RE = re.compile(r'[×xX]1(개|세트|박스|팩)')


def _strip_redundant_one_multiplier(s):
    """정식 옵션 목록에 "30봉×1세트"(기본 1개)와 "30봉×2세트"(2개)가 함께
    등록된 경우가 있는데, 실제 주문 원문은 기본값일 때 배수 표기 없이 그냥
    "30봉 세트"로만 오는 경우가 많다. "×1단위"를 "1"을 뗀 단위 단어만 남기고
    지워서("30봉×1세트" -> "30봉 세트") 이런 기본값 주문도 매치되게 한다.
    "×2세트"처럼 1이 아닌 배수는 그대로 둬 서로 다른 옵션으로 구분된다."""
    return _REDUNDANT_ONE_RE.sub(r' \1', s)


_PACKAGE_COUNT_RE = re.compile(r'(\d+)개입')


def _normalize_package_count(s):
    """정식 옵션 목록은 낱개 수를 "8입"으로 적지만, 원본 주문서는 "8개입"으로
    오는 경우가 있다("오리지날 8개입x2박스"). 둘 다 "8조각 들어있음"이라는
    같은 뜻이라 "N개입" -> "N입"로 통일해서 비교한다(배수 판단용 텍스트가
    아니라 옵션 식별용 텍스트에만 적용 — "45개입"을 구매 배수로 보지 않는
    _X_COUNT_RE의 제외 규칙과는 별개다)."""
    return _PACKAGE_COUNT_RE.sub(r'\1입', s)


_CONTAINER_UNIT_RE = re.compile(r'(\d)(박스|팩)')


def _normalize_container_words(s):
    """"박스"/"팩"/"세트"는 다 같은 "묶음" 뜻으로 마켓/협력사마다 표기가
    갈린다("30봉x2박스" 원문 vs 정식 옵션 목록의 "30봉x2세트"). 배수 판단·
    옵션 매칭에서는 이 단어 차이가 다른 옵션을 뜻하는 게 아니므로, 숫자
    뒤에 오는 "박스"/"팩"을 "세트"로 통일해서 비교한다(정식 옵션 목록
    토큰화할 때도 똑같이 적용해 양쪽이 같은 기준으로 맞는다). 배수 감지
    정규식(_BOX_TOTAL_RE, _X_COUNT_RE)은 이미 세 단어를 동등하게 취급하고
    있어 이 정규화와 일관된다."""
    return _CONTAINER_UNIT_RE.sub(r'\g<1>세트', s)


_CANDY_SYNONYM_RE = re.compile(r'사탕')


def _normalize_candy_word(s):
    """"사탕"(고유어)과 "캔디"(외래어)는 같은 뜻인데 마켓/참고표마다 표기가
    갈린다(이엑스 참고표엔 "카스텔리모 레몬사탕"처럼 "사탕"으로 등록돼
    있는데 실제 카카오 주문서엔 "레몬캔디"처럼 "캔디"로 오는 경우가 많음).
    "사탕" -> "캔디"로 통일해서 비교한다."""
    return _CANDY_SYNONYM_RE.sub('캔디', s)


_STICK_COUNT_RE = re.compile(r'(\d+)스틱')


def _normalize_stick_count(s):
    """정식 옵션 목록은 낱개 수를 "15포"로 적지만, 원본 주문서는 상품명에
    "10g*15스틱"처럼 "N스틱"으로 적는 경우가 있다(같은 낱개를 스틱형 포장
    이라고 부른 것뿐, 카탈로그에 "N스틱" 표기 상품은 없다). "N스틱" ->
    "N포"로 통일해서 비교한다."""
    return _STICK_COUNT_RE.sub(r'\1포', s)


_REDUNDANT_TRAILING_CONTAINER_RE = re.compile(
    r'([xX×]\s*\d+\s*(?:개(?!입)|세트|박스|팩))\s*(?:세트|박스|팩)'
)


def _strip_redundant_trailing_container(s):
    """"30포×2개 세트"처럼 "×N개"로 이미 배수를 밝힌 뒤에 "세트"/"박스"/
    "팩" 같은 묶음 단어가 한 번 더 따라붙는 경우가 있다(같은 뜻의 중복
    표현 — "2개짜리를 세트로 산다"는 말). 이 트레일링 단어까지 식별
    단어로 요구하면, 정식 옵션 목록엔 낱개 기준 이름만 있고 "세트"라는
    말이 안 붙어 있어서(예: "산골농장 하루한포 배도라지 스틱 30포")
    매칭이 실패한다. 배수는 이미 "×N개"로 판단되므로 뒤따르는 중복
    묶음 단어는 지운다."""
    return _REDUNDANT_TRAILING_CONTAINER_RE.sub(r'\1', s)


def _apply_option_synonyms(text, cfg):
    """원본 주문서 옵션이 줄임말/구어체로 오는 경우(예: "청포도"가 실제로는
    "청포도요거트" 맛을 가리키는데 정식 옵션 목록엔 "청포도"만 있는 맛은
    없고 "청포도요거트"만 있는 경우), 정식 옵션 텍스트로 미리 치환해서
    토큰 매칭이 되게 한다. 이미 정식 표현대로 온 경우(뒤에 no-op이 되도록
    부정형 전방탐색으로 이중 치환은 막는다. vendors.json의 "option_synonyms"
    ({줄임말: 정식표현}) 항목은 사용자에게 직접 확인받고 등록한 것만 넣는다."""
    synonyms = cfg.get("option_synonyms")
    if not synonyms or not text:
        return text
    for short, long_form in synonyms.items():
        if not short or short == long_form or not long_form.startswith(short):
            continue
        suffix = long_form[len(short):]
        text = re.sub(re.escape(short) + r'(?!' + re.escape(suffix) + r')', long_form, text)
    return text


def _normalize_for_match(text, cfg):
    """옵션 매칭에 쓰이는 원문 텍스트에 적용하는 정규화를 한데 모은다(동의어
    치환 -> 박스/팩/세트 통일 -> 개입/입 통일 순서). 정식 옵션 목록 쪽
    토큰화(_load_option_canon_table)에도 같은 통일 규칙(×1 배수 제거 제외)이
    적용돼 있어야 양쪽이 같은 기준으로 비교된다."""
    text = _apply_option_synonyms(text, cfg)
    text = _normalize_container_words(text)
    text = _normalize_package_count(text)
    text = _normalize_stick_count(text)
    text = _strip_redundant_trailing_container(text)
    text = _normalize_candy_word(text)
    return text


def _load_option_canon_table(vendor_name, cfg):
    if vendor_name not in _option_canon_tables:
        opt_file = cfg.get("option_canon_file")
        if not opt_file:
            _option_canon_tables[vendor_name] = None
        else:
            with open(BASE_DIR / opt_file, encoding="utf-8") as f:
                entries = json.load(f)
            items = [
                (
                    _tokenize(_strip_redundant_trailing_container(_normalize_stick_count(
                        _normalize_package_count(_normalize_container_words(
                            _strip_redundant_one_multiplier(e)
                        ))
                    ))),
                    e,
                )
                for e in entries
            ]
            items = [(t, e) for t, e in items if t]
            items.sort(key=lambda x: sum(len(t) for t in x[0]), reverse=True)
            _option_canon_tables[vendor_name] = items
    return _option_canon_tables[vendor_name]


# ---------------------------------------------------------------------------
# 수량 재산출 (가공 지침 1번): 정식 옵션 목록의 각 항목은 "낱개 단위" 하나를
# 뜻한다. 판매수량은 항상 "원본 수량 × 상품명/옵션에 명시된 배수"로 계산한다
# (원본 수량을 버리고 배수로 대체하는 게 아니라 곱한다). 배수 근거는 우선순위
# 순으로 셋:
#   1) "총 N단위" — 정식 옵션명 자체의 단위(예: "60포")를 기준으로 원문에
#      "총 120포"처럼 명시된 경우, 120//60=2배. 정식 옵션 매칭에 성공했을 때만
#      판단 가능(그 옵션의 고유 단위를 알아야 나눌 수 있으므로).
#   2) "(N박스)"/"(N세트)"/"(N팩)" — 괄호로 묶은 판매 단위 수. 상품명에 낱개
#      수("70g x 16개")가 같이 적혀 있어도 실제 판매 단위는 이 괄호 표기이므로
#      더 우선한다. 정식 옵션 매칭 여부와 무관하게(카탈로그에 없는 상품이라도)
#      원문에서 바로 판단할 수 있다.
#   3) "옵션 자체 수량 없이 그냥 x N개/세트/박스/팩" — "아마드티 얼그레이
#      20티백x6개"처럼 옵션명("20티백")에는 개수가 없고 상품명에 "xN개"만
#      곱셈으로 붙은 경우. 정식 옵션 매칭 여부와 무관하게 원문에서 판단한다.
#      "한입 허니 꽈배기 520gx3개"(옵션 "520g 3개")처럼 옵션 자체에 이미 그
#      수량이 포함된 것처럼 보여도, 상품명 쪽의 "x3개"는 여전히 "그 옵션을
#      3개 산다"는 뜻이므로 곱한다 — "3개"가 옵션명과 겹친다고 배수를 무시하지
#      않는다. 단, "45개입"처럼 "개" 바로 뒤에 "입"이 붙으면 그건 한 봉지 안에
#      든 낱개 수(포장 단위 설명)라 구매 배수가 아니므로 매칭에서 제외한다
#      (예: "15gx45개입, 675g" — 675g짜리 한 봉지 안에 15g들이 45개가 들어
#      있다는 뜻이지 45개를 주문한다는 뜻이 아니다).
# 위 셋 중 어느 것도 명시적으로 없으면 배수 1(=원본 수량 그대로)로 본다.
# ---------------------------------------------------------------------------
_UNIT_TOKEN_RE = re.compile(r'^(\d+)([가-힣a-zA-Z]+)$')
_TOTAL_RE = re.compile(r'총\s*(\d+)\s*([가-힣a-zA-Z]+)')
_BOX_TOTAL_RE = re.compile(r'\((\d+)\s*(?:박스|세트|팩)\)')
_X_COUNT_RE = re.compile(r'[xX×]\s*(\d+)\s*(?:개(?!입)|세트|박스|팩)')
_EMBEDDED_MULTIPLIER_RE = re.compile(r'[×xX]\s*([2-9]\d*)\s*(?:개(?!입)|세트|박스|팩)')


def _has_embedded_multiplier(matched_option):
    """정식 옵션명 자체가 이미 배수/여러 구성품을 포함한 하나의 완결된 SKU라
    상품명의 "총 N" 같은 문구를 판매수량에 추가로 곱하면 안 되는 경우를
    본다. 두 가지:

    1) "×2세트"처럼 1보다 큰 배수가 이미 포함된 경우(예: "요거트 30봉×1세트"와
       "요거트 30봉×2세트"가 각각 별도로 정식 옵션 목록에 등록된 경우 —
       후자는 "2세트짜리 묶음" 자체가 하나의 선택지다).
    2) "후르츠 30봉+요거트 30봉+프리미엄 30봉 세트"처럼 "+"로 서로 다른 맛을
       묶은 조합 SKU인 경우 — 상품명의 "총 90봉"은 이미 그 세 가지 맛
       30봉씩을 합친 값을 설명하는 것이지, 이 조합을 3번 사라는 뜻이 아니다.

    두 경우 다, 옵션을 몇 번 골랐는지(판매수량)와 옵션 자체에 이미 포함된
    배수/구성을 이중으로 곱하면 안 된다 — 상품명에 그 묶음의 총 개수를
    설명하는 문구가 같이 있어도, 그건 이미 옵션에 포함된 내용을 설명하는
    것이지 추가 구매 배수가 아니다."""
    return bool(_EMBEDDED_MULTIPLIER_RE.search(matched_option)) or "+" in matched_option


def _detect_multiplier(matched_tokens, raw_text):
    """matched_tokens: 정식 옵션명에서 뽑은 필수 토큰 집합(그중 "60포"처럼
    숫자+단위인 것만 배수 판단에 쓴다). raw_text: 원본 상품명+옵션 원문(공백
    유지). "총 N단위"처럼 명시적인 근거가 있을 때만 배수를 반환한다 — 그 외에는
    1을 반환하고(=재계산 안 함), 호출부에서 _detect_general_multiplier로
    이어서 판단한다."""
    unit_tokens = []
    for tok in matched_tokens:
        m = _UNIT_TOKEN_RE.match(tok)
        if m:
            unit_tokens.append((int(m.group(1)), m.group(2)))
    if not unit_tokens:
        return 1

    for total_str, unit in _TOTAL_RE.findall(raw_text):
        total = int(total_str)
        for base, base_unit in unit_tokens:
            if base_unit == unit and base and total % base == 0:
                mult = total // base
                if mult >= 1:
                    return mult

    return 1


def _detect_general_multiplier(raw_text):
    """정식 옵션 매칭 성공 여부와 무관하게 원문 자체만으로 판단 가능한 배수
    (위 우선순위 2, 3번). 괄호 묶음 수를 먼저 보되, "1박스"처럼 1묶음이라고만
    적힌 경우는 배수라고 할 게 없는 경우라("33g, 10개 (1박스)"처럼 낱개 수
    "10개"가 진짜 수량이고 "1박스"는 그 낱개들이 박스 하나에 들었다는
    설명일 뿐인 경우가 있다) 그 경우엔 배수 후보로 치지 않고 다음(xN개
    곱셈 표기)으로 넘어간다."""
    m = _BOX_TOTAL_RE.search(raw_text)
    if m:
        mult = int(m.group(1))
        if mult > 1:
            return mult
    m = _X_COUNT_RE.search(raw_text)
    if m:
        mult = int(m.group(1))
        if mult >= 1:
            return mult
    return 1


_BARE_COUNT_RE = re.compile(r'(?<![×xX])(\d+)개(?!입)')


def _detect_bare_count(text):
    """"x"/"×" 곱셈 표기 없이 그냥 "N개"라고만 적힌 낱개 수를 찾는다(예:
    "브라카 커피 비스킷 150g, 4개" — 옵션 필드가 아예 없어서 상품명 자체가
    이 리스팅이 몇 개들이인지 말해주는 유일한 정보인 경우에 쓴다). 못
    찾으면 None을 반환한다(1을 반환하면 "명시적으로 1개"인지 "아무 근거
    없음"인지 구분이 안 되므로)."""
    m = _BARE_COUNT_RE.search(text)
    return int(m.group(1)) if m else None


# ---------------------------------------------------------------------------
# 한성로직스: 상품명/옵션명을 항상 옵션 텍스트로 통일하고, "×N단위" 배수는
# 표시상 "×1단위"로 고정한 뒤 그 배수를 판매수량에 반영한다(사용자가 직접
# 확인해준 지침 — 2026-09-11: 예) "선택: 예향 한입방울떡 1kg×2개" ->
# "선택: 예향 한입방울떡 1kg×1개", 수량 1 -> 2).
# ---------------------------------------------------------------------------
_X_MULTIPLIER_RE = re.compile(r'[×xX](\d+)(개|봉|입|포|세트|박스|팩)')


def _normalize_x_multiplier(vendor_name, cfg, rec):
    if not cfg.get("normalize_x_multiplier"):
        return rec
    option_text = rec.get("option") or ""
    if not option_text.strip():
        return rec
    new_rec = dict(rec)
    m = _X_MULTIPLIER_RE.search(option_text)
    if m:
        count = int(m.group(1))
        unit = m.group(2)
        option_text = option_text[:m.start()] + f"×1{unit}" + option_text[m.end():]
        new_rec["quantity"] = (rec.get("quantity") or 1) * count
    new_rec["option"] = option_text
    new_rec["product_name"] = option_text
    return new_rec


# ---------------------------------------------------------------------------
# 같은 단위가 "+"로 반복된 표기 정리: "60포+60포"처럼 완전히 동일한
# 숫자+단위가 "+"로 반복되면, 이건 서로 다른 구성품의 조합이 아니라 같은
# 상품을 여러 번 산다는 뜻이다(가공 지침 1번의 "총 120포" 예시 — 60포짜리를
# 2번). 정식 옵션 목록 매칭 여부와 무관하게 원문 자체에서 판단 가능하며,
# 매칭시켜서 문구를 바꿀 필요도 없다 — 반복된 부분만 하나로 줄이고 나머지
# 원문(예: "(원통)")은 그대로 둔 채 판매수량만 반복 횟수만큼 곱한다. 정식
# 옵션 매칭을 거치지 않으므로 "(원통)"처럼 카탈로그에 없는 포장 설명이
# 붙어 있어도 원문 그대로 정확히 보존된다(카탈로그의 "60포" 단품과는 다른
# 포장일 수 있어 억지로 정규화하면 안 된다).
# ---------------------------------------------------------------------------
_REPEATED_UNIT_RE = re.compile(r'(\d+[가-힣]+)(\+\1)+')


def _collapse_repeated_unit(rec):
    option_text = rec.get("option") or ""
    m = _REPEATED_UNIT_RE.search(option_text)
    if not m:
        return rec
    repeated_span = m.group(0)
    unit = m.group(1)
    count = repeated_span.count("+") + 1
    new_rec = dict(rec)
    new_rec["option"] = option_text.replace(repeated_span, unit, 1)
    new_rec["quantity"] = (rec.get("quantity") or 1) * count
    # 배수를 이미 여기서 확정했다는 표시. 상품명 쪽에 "(총 120포)"처럼 같은
    # 총량을 알려주는 문구가 남아 있으면 뒤의 _apply_option_recalc가 그걸
    # 보고 또 배수를 곱해버리는 사고(2배 -> 4배)가 나므로, 이 표시가 있으면
    # 배수 재계산을 건너뛰게 한다.
    new_rec["_repeated_unit_collapsed"] = True
    return new_rec


# ---------------------------------------------------------------------------
# 다중 선택 행 분리 (가공 지침 2번): 고객이 "(택N)"으로 서로 다른 옵션 N개를
# 선택하면, 원본 옵션 필드에 그 N개 선택 내역이 구분자(콤마/슬래시/"선택N:"
# 표기 등)로 나뉘어 실제로 들어있다. 그 구분된 항목들이 각각 정식 옵션
# 목록과 confident하게 매칭되고 서로 다른 항목이면, 그 개수만큼 행을 나눈다.
# 구분한 항목 중 하나라도 매칭이 안 되거나 같은 옵션으로 겹치면(=진짜 여러
# 선택인지 확신할 수 없으면) 나누지 않고 원본 그대로 둔다.
# ---------------------------------------------------------------------------
_OPTION_SEGMENT_RE = re.compile(r'선택\s*\d*\s*[:.]|[,/\n]')


def _split_option_segments(option_text):
    if not option_text:
        return []
    parts = _OPTION_SEGMENT_RE.split(option_text)
    return [p.strip() for p in parts if p and p.strip()]


def _expand_multi_select(vendor_name, cfg, rec):
    """옵션 필드가 서로 다른 정식 옵션 2개 이상을 담고 있으면 그만큼 행을
    나눠 반환한다(각 행은 낱개 1개씩으로 본다 — "총 수량"이 아니라 "고객이
    고른 서로 다른 항목의 개수"이므로). 분리 대상이 아니거나 애매하면
    [rec] 그대로(리스트 하나) 반환한다."""
    table = _load_option_canon_table(vendor_name, cfg)
    if not table:
        return [rec]
    segments = _split_option_segments(rec.get("option"))
    if len(segments) < 2:
        return [rec]

    matched = []
    for seg in segments:
        result = _best_token_match_with_tokens(seg, table)
        if not result:
            return [rec]
        matched.append(result[1])
    if len(set(matched)) != len(matched):
        return [rec]

    out = []
    for opt in matched:
        new_rec = dict(rec)
        new_rec["option"] = opt
        new_rec["quantity"] = 1
        out.append(new_rec)
    return out


_OPTION_BOILERPLATE_TOKENS = {"선택", "옵션", "원통", "랜덤발송", "선물용"}


def _find_option_match(table, option_text, combined_text):
    """옵션 텍스트만으로 먼저 매칭을 시도하고, 못 찾으면 상품명을 더한 전체
    텍스트로 넘어간다(식별 단어 검증 포함) — _apply_option_recalc와
    _expand_plus_combo가 공유하는 매칭 로직. 매칭되면 (tokens, matched_option,
    실제 매칭에 쓰인 텍스트)를, 못 찾으면 None을 반환한다."""
    result = _best_token_match_with_tokens(option_text, table) if option_text else None
    if result:
        return result[0], result[1], option_text
    option_identity_tokens = {
        t for t in _tokenize(option_text)
        if not t[0].isdigit() and t not in _OPTION_BOILERPLATE_TOKENS
    }
    candidate = _best_validated_match(combined_text, table, option_identity_tokens)
    if candidate:
        return candidate[0], candidate[1], combined_text
    candidate = _best_relaxed_match(option_text, table, option_identity_tokens)
    if candidate:
        return candidate[0], candidate[1], option_text
    return None


def _multiplier_for_match(matched_option, tokens, match_text, general_text):
    """정식 옵션 매칭에 성공했을 때의 배수를 판단한다(옵션 자체에 이미
    배수가 포함돼 있으면 1, 아니면 "총 N" -> 일반 배수 순으로 확인)."""
    if _has_embedded_multiplier(matched_option):
        return 1
    multiplier = _detect_multiplier(tokens, match_text)
    if multiplier == 1:
        multiplier = _detect_general_multiplier(general_text)
    return multiplier


def _expand_box_count_bundle(vendor_name, cfg, rec):
    """일부 상품(예: 이뮨 "퍼펙트이뮨")은 옵션 텍스트로 맛/종류를 매칭하는
    게 아니라, "몇 박스를 샀는지"에 따라 본품 + 부속품(쇼핑백 종류는
    박스 수에 따라 다름, 특정 수량 이상이면 단상자 추가, 옵션에 증정
    문구가 있으면 증정품 추가)으로 항상 고정된 구성으로 나뉘어야 하는
    경우가 있다(사용자가 직접 확인해준 지침 — 2026-09-11). cfg["box_
    count_bundle"]에 그런 규칙이 등록돼 있으면 상품명+옵션에서 박스
    개수를 판단해 구성요소별로 행을 나눈다. 해당 없으면 [rec] 그대로
    반환한다."""
    rule = cfg.get("box_count_bundle")
    if not rule:
        return [rec]
    product_name = rec.get("product_name") or ""
    option_text = rec.get("option") or ""
    combined = _normalize_for_match(f"{product_name} {option_text}", cfg)
    box_count = _detect_general_multiplier(combined)
    if box_count < 1:
        box_count = 1

    items = [(rule["unit_product"], box_count)]
    items.append((rule["bag_single"] if box_count == 1 else rule["bag_multi"], 1))
    if box_count >= rule.get("box_from_count", 10 ** 9):
        items.append((rule["box_item"], 1))
    gift_trigger = rule.get("gift_trigger")
    if gift_trigger and gift_trigger in option_text:
        items.append((rule["gift_item"], 1))

    out = []
    for name, qty in items:
        new_rec = dict(rec)
        new_rec["product_name"] = name
        new_rec["option"] = ""
        new_rec["quantity"] = qty
        out.append(new_rec)
    return out


def _expand_variety_set(vendor_name, cfg, rec):
    """일부 상품은 정식 옵션 목록에 "여러 캐릭터/맛을 합친" 콤보 항목이
    있어도 실제로는 항상 "각각 하나씩"을 뜻해서 그 콤보 항목과 매칭하면
    안 되고, vendors.json에 등록해둔 개별 구성 요소들로 무조건 나눠야
    하는 경우가 있다(예: 먼작귀 보틀피규어 "3종 세트" — 사용자가 직접
    확인해준 특이 케이스). cfg["variety_set_groups"]에 그런 상품군이
    등록돼 있고 상품명이 그 그룹의 product_contains 키워드를 전부
    포함하면, 원문에서 판단되는 배수(세트 수)만큼 구성 요소 각각에
    적용해 그 개수만큼 행을 나눈다. 해당 없으면 [rec] 그대로 반환한다."""
    groups = cfg.get("variety_set_groups")
    if not groups:
        return [rec]
    product_name = rec.get("product_name") or ""
    option_text = rec.get("option") or ""
    for group in groups:
        if not all(kw in product_name for kw in group["product_contains"]):
            continue
        combined = _normalize_for_match(f"{product_name} {option_text}", cfg)
        multiplier = _detect_general_multiplier(combined)
        quantity = (rec.get("quantity") or 1) * multiplier
        out = []
        for component in group["components"]:
            new_rec = dict(rec)
            new_rec["option"] = component
            new_rec["quantity"] = quantity
            out.append(new_rec)
        return out
    return [rec]


def _match_plus_segments(table, cfg, product_name, segments, clean_others):
    """세그먼트들을 각각 정식 옵션과 매칭해본다. clean_others=True면 각
    세그먼트를 매칭할 때 "다른" 세그먼트들의 식별 단어를 상품명에서 지운
    텍스트를 쓴다(예: "레몬 200g+모로오렌지 200g"에서 "레몬"을 매칭할 때
    상품명의 "모로오렌지" 단어를 지워서 "레몬"이 "캔디"와 바로 붙게 만듦
    — 상품명이 "레몬/모로오렌지 캔디사탕"처럼 두 맛을 나란히 소개하는
    문구라 원래는 "모로오렌지"가 "레몬"과 "캔디" 사이에 끼어들어 매칭이
    안 되는 문제를 해결). 전부 매칭되면 [(matched_option, multiplier), ...]
    를, 하나라도 실패하면 None을 반환한다."""
    matched = []
    for i, seg in enumerate(segments):
        seg_combined_name = product_name
        if clean_others:
            other_words = set()
            for j, other_seg in enumerate(segments):
                if j == i:
                    continue
                other_words |= {t for t in _tokenize(other_seg) if not t[0].isdigit()}
            for w in other_words:
                seg_combined_name = seg_combined_name.replace(w, "")
        seg_norm = _normalize_for_match(seg, cfg)
        seg_combined = _normalize_for_match(f"{seg_combined_name} {seg}", cfg)
        found = _find_option_match(table, seg_norm, seg_combined)
        if not found:
            return None
        tokens, matched_option, match_text = found
        multiplier = _multiplier_for_match(matched_option, tokens, match_text, seg_norm)
        # "레몬 2개+오렌지 2개"처럼 세그먼트 자체에 "×" 없는 맨 "N개"로 그
        # 조각의 개수가 적혀 있는 경우("×N개" 형태만 보는 _detect_general_
        # multiplier로는 못 잡음) 그 조각의 판매수량으로 그대로 쓴다.
        if multiplier == 1:
            bare_count = _detect_bare_count(seg)
            if bare_count:
                multiplier = bare_count
        matched.append((matched_option, multiplier))
    return matched


def _expand_plus_combo(vendor_name, cfg, rec):
    """"+"로 여러 항목이 묶인 옵션을 처리한다. 먼저 전체 옵션 텍스트가
    그 자체로 하나의 조합 SKU로 정식 옵션 목록에 등록돼 있는지 본다(예:
    "후르츠 30봉+요거트 30봉+프리미엄 30봉 세트" — 이런 조합 자체가 하나의
    상품으로 등록돼 있으면 한 줄로 그대로 두고 _apply_option_recalc가
    처리하게 한다). 등록된 조합이 없으면 "+"로 나눠 각 조각을 독립적으로
    매칭해본다(예: "오리지날 8개입x2박스+초코 8개입x2박스" -> "추억의도나스
    오리지날 8입"/"추억의도나스 초코 8입"이 각각 따로 등록돼 있는 경우).
    이게 안 되면(상품명에 다른 조각의 맛 이름이 끼어들어 있어 세그먼트
    매칭이 실패하는 경우) 다른 조각들의 식별 단어를 상품명에서 지운
    텍스트로 다시 각 조각을 매칭해본다("레몬 200g+모로오렌지 200g" 같은
    같은 상품의 서로 다른 맛 조합). 둘 중 하나라도 전체 조각이 매칭되고
    서로 다른 정식 옵션이면(같은 옵션으로 겹치지 않으면) 그 개수만큼 행을
    나누고, 각 행의 수량은 그 조각 자체의 배수(임베디드 배수 또는 "xN개/
    세트/박스/팩")를 반영한다. 둘 다 안 되면 원본 그대로 둔다(안전하게
    사람이 확인하도록)."""
    table = _load_option_canon_table(vendor_name, cfg)
    if not table:
        return [rec]
    option_text_raw = (rec.get("option") or "").strip()
    if "+" not in option_text_raw:
        return [rec]

    product_name = rec.get("product_name") or ""
    whole_option = _normalize_for_match(option_text_raw, cfg)
    whole_combined = _normalize_for_match(f"{product_name} {option_text_raw}", cfg)
    if _find_option_match(table, whole_option, whole_combined):
        return [rec]

    segments = [s.strip() for s in option_text_raw.split("+") if s.strip()]
    if len(segments) < 2:
        return [rec]

    matched = _match_plus_segments(table, cfg, product_name, segments, clean_others=False)
    if matched is None:
        matched = _match_plus_segments(table, cfg, product_name, segments, clean_others=True)

    if matched is not None:
        matched_options = [m[0] for m in matched]
        if len(set(matched_options)) == len(matched_options):
            out = []
            for matched_option, multiplier in matched:
                new_rec = dict(rec)
                new_rec["option"] = matched_option
                new_rec["quantity"] = (rec.get("quantity") or 1) * multiplier
                out.append(new_rec)
            return out

    return [rec]


def _apply_option_recalc(vendor_name, cfg, rec):
    """옵션 정규화 + 필요 시 수량 재산출을 반영한 새 rec를 반환한다(원본은
    건드리지 않음). 정식 옵션 매칭에 성공하면 옵션명을 그걸로 바꾸고 "총 N"
    배수까지 확인한다. 매칭에 실패해도(카탈로그에 없는 상품이라도) 원문에
    괄호 묶음 수나 "xN개" 표기가 있으면 배수만 반영한다(옵션명은 원본 그대로
    둔다) — 예: "추억의 도나스 70g x 16개 (2박스)"는 정식 옵션 목록에 없지만
    "(2박스)"는 그대로 판단 가능하다.

    옵션 매칭은 옵션 텍스트만으로 먼저 시도하고, 거기서 못 찾을 때만 상품명을
    더한 전체 텍스트로 넘어간다. "OO 모음전 / 후르츠,요거트,치즈퀴노아 등"처럼
    상품명 자체가 판매 중인 여러 맛을 소개하는 문구인 경우가 있어서, 고객이
    실제로 고른 건 "요거트" 단품인데 상품명에 다른 맛 이름들이 같이 있다는
    이유만으로 더 구체적인(토큰이 많은) 조합 옵션("후르츠+요거트 세트" 등)에
    잘못 매칭될 수 있기 때문이다 — 옵션 텍스트 자체만으로 이미 확정 가능한
    경우에는 상품명을 끌어들이지 않는다. 옵션 텍스트만으로 못 찾을 때만(예:
    옵션에 브랜드/상품 식별 정보가 아예 없는 이플코리아 케이스) 상품명을 더한
    전체 텍스트로 재시도한다.

    상품명을 더한 전체 텍스트로 찾은 후보는 한 번 더 검증한다: 옵션 텍스트
    자체에 있는 "식별 단어"(숫자로 시작하지 않는 한글 토큰 — "검은콩오곡"
    같은 맛/품목 이름. "20개입"처럼 숫자로 시작하는 수량 표기 토큰은 애초에
    상품명 쪽에 있는 게 정상이라 제외하고, "선택"/"옵션"처럼 마켓 export가
    옵션 필드 앞에 습관적으로 붙이는 상투어도 제외 — 정식 옵션 목록 항목이
    "선택:"으로 시작하지 않는 협력사가 많아서, 이 단어를 식별 단어로 치면
    정상적인 매칭까지 다 막혀버린다)가 후보의 필수 토큰에 전부 포함돼 있어야
    그 매칭을 받아들인다. 안 그러면, 고객이 실제로 고른 맛이 정식 옵션
    목록에 아예 없는 경우(예: "검은콩오곡")에 상품명의 다른 맛 이름들이
    우연히 다 모여서 전혀 다른 조합 옵션에 매칭돼버리는 사고가 난다 — 이럴
    땐 억지로 맞추지 말고 원본 그대로 두는 게 안전하다."""
    new_rec = dict(rec)
    already_resolved = new_rec.pop("_repeated_unit_collapsed", False)

    table = _load_option_canon_table(vendor_name, cfg)
    if not table:
        return new_rec
    option_text = _normalize_for_match((rec.get("option") or "").strip(), cfg)
    combined_text = _normalize_for_match(
        f"{rec.get('product_name') or ''} {rec.get('option') or ''}", cfg
    )
    if not combined_text.strip():
        return new_rec

    found = _find_option_match(table, option_text, combined_text)

    if found:
        tokens, matched_option, match_text = found
        new_rec["option"] = matched_option
        multiplier = 1 if already_resolved else _multiplier_for_match(
            matched_option, tokens, match_text, combined_text
        )
    else:
        multiplier = 1 if already_resolved else _detect_general_multiplier(combined_text)

    # 옵션 필드가 아예 비어 있으면 상품명이 그 리스팅의 개수를 말해주는
    # 유일한 정보이므로, 다른 배수 근거가 없을 때 상품명의 "N개"(곱셈
    # 표기 없는 낱개 수)를 그대로 판매수량으로 쓴다.
    if multiplier == 1 and not already_resolved and not option_text.strip():
        bare_count = _detect_bare_count(rec.get("product_name") or "")
        if bare_count:
            multiplier = bare_count

    if multiplier > 1:
        new_rec["quantity"] = (rec.get("quantity") or 1) * multiplier

    return new_rec


def _capture_row_style(ws, row_idx, max_col):
    styles = {}
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row_idx, column=c)
        styles[c] = {
            "font": copy(cell.font),
            "fill": copy(cell.fill),
            "border": copy(cell.border),
            "alignment": copy(cell.alignment),
            "number_format": cell.number_format,
        }
    return styles


_PHONE_DIGITS_RE = re.compile(r'\D+')


def _format_phone_hyphenated(phone):
    """"01012345678"처럼 하이픈 없이 오는 전화번호를 "010-1234-5678"
    (일반적인 xxx-xxxx-xxxx 표기)로 바꾼다. 11자리(휴대폰)는 3-4-4,
    10자리(구형 번호/지역번호 포함 유선)는 3-3-4로 나눈다. 그 외
    자릿수는 형식을 판단할 수 없어 원본 그대로 둔다."""
    digits = _PHONE_DIGITS_RE.sub("", phone or "")
    if len(digits) == 11:
        return f"{digits[:3]}-{digits[3:7]}-{digits[7:]}"
    if len(digits) == 10:
        return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}"
    return phone or None


def _cell_value_for(kind_spec, rec, seq, is_first_of_order, vendor_name=None, cfg=None, group_seq=None):
    kind = kind_spec[0]
    if kind == "seq":
        return seq
    if kind == "year":
        return f"{rec['order_date'].year}년" if rec.get("order_date") else None
    if kind == "month":
        return rec["order_date"].month if rec.get("order_date") else None
    if kind == "day":
        return rec["order_date"].day if rec.get("order_date") else None
    if kind == "static":
        return kind_spec[1]
    if kind == "combine":
        parts = [p for p in [rec.get("product_name"), rec.get("option")] if p]
        return " ".join(parts)
    if kind == "shipment_flag":
        return 1 if is_first_of_order else None
    if kind == "consolidated":
        return kind_spec[1] if is_first_of_order else None
    if kind == "group_seq":
        return group_seq if is_first_of_order else None
    if kind == "field":
        if kind_spec[1] == "receiver_phone_hyphenated":
            return _format_phone_hyphenated(rec.get("receiver_phone"))
        return rec.get(kind_spec[1]) or None
    if kind == "code":
        return _code_for(vendor_name, cfg, rec)
    if kind == "name_pair":
        match = _name_pair_match(vendor_name, cfg, rec)
        if match:
            return match[kind_spec[1]]
        if kind_spec[1] == "a":
            return rec.get("product_name") or None
        return rec.get("option") or rec.get("product_name") or None
    return None


def _find_col_by_kind(col_map, predicate):
    for c, spec in col_map.items():
        if predicate(spec):
            return c
    return None


def _dup_key(rec):
    name = (rec.get("receiver_name") or "").strip().casefold()
    addr = (rec.get("address") or "").strip().casefold()
    if not name or not addr:
        return None
    return (name, addr)


def _highlight_fill_for(rec, dup_counts, name_pair_unmatched=False):
    if name_pair_unmatched:
        return UNMATCHED_NAME_FILL
    is_qty = (rec.get("quantity") or 1) > 1
    key = _dup_key(rec)
    is_dup = key is not None and dup_counts[key] >= 2
    if is_qty and is_dup:
        return BOTH_FILL
    if is_qty:
        return QTY_FILL
    if is_dup:
        return DUP_FILL
    return None


def write_vendor_file(vendor_name, rows, out_path):
    """vendor_name 협력사 양식 템플릿을 복사해 rows(표준 필드 dict 리스트)를 채운다."""
    cfg = VENDORS[vendor_name]
    template_path = BASE_DIR / cfg["template_file"]
    wb = openpyxl.load_workbook(template_path)
    ws = wb.worksheets[0]

    col_map, dup_cols = _build_col_map(ws, cfg)
    for c in dup_cols:
        ws.column_dimensions[get_column_letter(c)].hidden = True
    data_start = cfg["data_start_row"]
    styles = _capture_row_style(ws, data_start, ws.max_column)

    # 옵션 정규화·수량 재산출(1번)/다중 선택 행 분리(2번) 지침 반영: 먼저 한 줄이
    # 서로 다른 정식 옵션 여러 개로 쪼개져야 하는지 보고(분리되면 그걸로 확정),
    # 등록된 "무조건 개별 구성요소로 나누는" 상품군인지 보고(해당하면 확정),
    # 아니면 "+"로 묶인 서로 다른 상품 조합인지 보고(이것도 분리되면 확정),
    # 그것도 아니면 같은 옵션의 배수인지(수량 재산출)만 확인한다.
    expanded_rows = []
    for rec in rows:
        rec = _normalize_x_multiplier(vendor_name, cfg, rec)
        bundle_split = _expand_box_count_bundle(vendor_name, cfg, rec)
        if len(bundle_split) > 1:
            expanded_rows.extend(bundle_split)
            continue
        rec = _collapse_repeated_unit(bundle_split[0])
        split_recs = _expand_multi_select(vendor_name, cfg, rec)
        if len(split_recs) > 1:
            expanded_rows.extend(split_recs)
            continue
        variety_split = _expand_variety_set(vendor_name, cfg, split_recs[0])
        if len(variety_split) > 1:
            expanded_rows.extend(variety_split)
            continue
        plus_split = _expand_plus_combo(vendor_name, cfg, variety_split[0])
        if len(plus_split) > 1:
            expanded_rows.extend(plus_split)
        else:
            expanded_rows.append(_apply_option_recalc(vendor_name, cfg, plus_split[0]))

    # 같은 수령인+주소로 가는 "서로 다른 주문"이 몇 건인지 세야 한다. 한 주문이
    # 여러 행으로 나뉘는 경우(옵션 분리, 이뮨 같은 구성품 묶음 등) 그 행들은
    # 전부 같은 주문번호를 공유하므로, 행 수가 아니라 주문번호 개수로 세지
    # 않으면 한 주문의 부속 행들끼리만 있어도 항상 "중복"으로 잘못 표시된다.
    dup_key_orders = {}
    for r in expanded_rows:
        key = _dup_key(r)
        if key is None:
            continue
        dup_key_orders.setdefault(key, set()).add(r.get("order_id"))
    dup_counts = {k: len(v) for k, v in dup_key_orders.items()}
    highlight_counts = {"quantity": 0, "duplicate_address": 0, "both": 0, "name_pair_unmatched": 0}
    has_name_pair = bool(cfg.get("name_pair_lookup_file"))

    seen_orders = set()
    group_seq = 0
    for i, rec in enumerate(expanded_rows):
        r = data_start + i
        order_id = rec.get("order_id")
        is_first = order_id not in seen_orders if order_id is not None else True
        if order_id is not None:
            seen_orders.add(order_id)
        if is_first:
            group_seq += 1

        for c, spec in col_map.items():
            value = _cell_value_for(spec, rec, i + 1, is_first, vendor_name, cfg, group_seq)
            cell = ws.cell(row=r, column=c)

            if spec[0] == "field" and spec[1] == "order_date" and value is not None:
                cell.value = value
                cell.number_format = "yyyy-mm-dd"
            else:
                cell.value = value

            st = styles.get(c)
            if st:
                cell.font = st["font"]
                cell.fill = st["fill"]
                cell.border = st["border"]
                cell.alignment = st["alignment"]
                if not (spec[0] == "field" and spec[1] == "order_date"):
                    cell.number_format = st["number_format"]

        name_pair_unmatched = has_name_pair and _name_pair_match(vendor_name, cfg, rec) is None
        fill = _highlight_fill_for(rec, dup_counts, name_pair_unmatched)
        if fill is UNMATCHED_NAME_FILL:
            highlight_counts["name_pair_unmatched"] += 1
        elif fill is QTY_FILL:
            highlight_counts["quantity"] += 1
        elif fill is DUP_FILL:
            highlight_counts["duplicate_address"] += 1
        elif fill is BOTH_FILL:
            highlight_counts["both"] += 1
        if fill is not None:
            for c in range(1, ws.max_column + 1):
                ws.cell(row=r, column=c).fill = fill

    if cfg.get("template_family") == "consignment_v1" and "summary_row" in cfg:
        sr = cfg["summary_row"]
        shipment_col = _find_col_by_kind(col_map, lambda s: s[0] == "shipment_flag")
        qty_col = _find_col_by_kind(col_map, lambda s: s == ("field", "quantity"))
        if shipment_col:
            letter = ws.cell(row=sr, column=shipment_col).column_letter
            ws.cell(row=sr, column=shipment_col).value = f"=SUM({letter}{data_start}:{letter}1048576)"
        if qty_col:
            letter = ws.cell(row=sr, column=qty_col).column_letter
            ws.cell(row=sr, column=qty_col).value = f"=SUM({letter}{data_start}:{letter}1048576)"

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path, highlight_counts


_REVIEW_HEADERS = [
    ("_source_file", "원본 파일"),
    ("_source_row", "원본 행"),
    ("order_id", "주문번호"),
    ("order_date", "주문일자"),
    ("product_name", "상품명"),
    ("option", "옵션"),
    ("quantity", "수량"),
    ("receiver_name", "수령인명"),
    ("receiver_phone", "수령인 연락처"),
    ("address", "배송지 주소"),
    ("zipcode", "우편번호"),
    ("shop_name", "쇼핑몰명"),
]


def write_review_file(unclassified, ambiguous, out_path):
    """협력사 매칭이 안 되었거나(미분류) 여러 협력사에 동시에 걸린(중복매칭) 건을
    수동 확인용 엑셀로 저장한다."""
    wb = openpyxl.Workbook()
    header_font = Font(name="맑은 고딕", size=10, bold=True)

    def write_sheet(ws, title, records, note_field=None):
        ws.title = title
        headers = [h for _, h in _REVIEW_HEADERS] + (["매칭 후보"] if note_field else [])
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = header_font
        for r, rec in enumerate(records, start=2):
            for c, (field, _) in enumerate(_REVIEW_HEADERS, start=1):
                ws.cell(row=r, column=c, value=rec.get(field))
            if note_field:
                ws.cell(row=r, column=len(_REVIEW_HEADERS) + 1, value=", ".join(rec.get(note_field, [])))

    ws1 = wb.active
    write_sheet(ws1, "미분류", unclassified)
    ws2 = wb.create_sheet("중복매칭")
    write_sheet(ws2, "중복매칭", ambiguous, note_field="_ambiguous_matches")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
