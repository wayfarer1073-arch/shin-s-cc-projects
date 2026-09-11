"""기간을 지정하면 보관된 주문 원본(data/orders/)에서 해당 브랜드의 모든 주문
(어느 협력사 것이든)을 모아 정산 리포트 공용 양식(templates/정산_*.xlsx)에 채운다.

배송처 컬럼은 주문마다 실제로 분류된 협력사로 채운다. 지급매입처 컬럼은
브랜드 상품 카탈로그(data/reference/catalog_*.json)에 그 상품의 pay_vendor가
등록돼 있으면 그걸 쓰고(같은 협력사라도 실제 매입 대금을 지급하는 곳이
따로 있는 경우가 있음 — 예: 투데이넛 상품의 매입처는 해맑음푸드), 없으면
배송처와 같다고 보고 분류된 협력사로 채운다. 관리상품명/매입단가도 같은
카탈로그에서 상품명+옵션이 정확히 일치하는 경우에만 채운다 (정확도가 낮은
매칭은 하지 않음).
그 외 우리 데이터로 확인 불가능한 값(카탈로그에 없는 상품의 원가/이익 등)은
공란으로 남긴다 — 사람이 템플릿에 같이 들어있는 참고 시트를 보고 채우게 된다.
"""
import json
import tempfile
from datetime import date, datetime
from pathlib import Path

import openpyxl

from src.archive import load_orders, ARCHIVE_DIR
from src.protect import protect_xlsx

BASE_DIR = Path(__file__).resolve().parent.parent

with open(BASE_DIR / "config" / "settlement_reports.json", encoding="utf-8") as f:
    _FULL_CFG = json.load(f)
    _CFG = _FULL_CFG["brands"]
    EXPORT_PASSWORD = _FULL_CFG.get("export_password")

# 열 번호(1-based) -> 의미. templates/정산_JM.xlsx (TEMPLATE(JM)) 기준.
COL = {
    "no": 1, "year": 2, "recv_month": 3, "recv_day": 4,
    "sales_channel": 7, "revenue_channel": 8, "pay_vendor": 9,
    "return_office": 10, "ship_office": 11, "order_id": 12,
    "receiver_name": 13, "receiver_phone": 14, "address": 15, "zipcode": 16,
    "product_name": 17, "option": 18, "managed_name": 19,
    "shipment_count": 20, "quantity": 21, "revenue": 22, "settlement": 23,
    "unit_cost": 24, "cost_amount": 25, "profit": 26, "margin": 27,
    "sender_name": 28, "sender_phone": 29, "sender_address": 30,
    "delivery_message": 31, "note": 42,
}


def _parse_date(s):
    return datetime.strptime(s, "%Y-%m-%d").date() if isinstance(s, str) else s


def _load_catalog_lookup(catalog_file):
    with open(BASE_DIR / catalog_file, encoding="utf-8") as f:
        entries = json.load(f)
    return {(e["order_product_name"], e["option"]): e for e in entries}


def generate_settlement_report(
    brand, start_date, end_date, out_path, archive_dir=ARCHIVE_DIR,
    protect=False, password=None,
):
    """start_date, end_date: "YYYY-MM-DD" 문자열. 반환: (out_path, 채운 건수).

    brand의 모든 주문(협력사 무관)을 기간으로 모아 정산 리포트를 채운다.
    protect=True로 켜면 결과 파일에 password(또는 config의 export_password)로
    열기암호를 걸어 저장한다. 기본은 암호 없이 생성한다."""
    if brand not in _CFG:
        raise ValueError(f"'{brand}' 브랜드의 정산 리포트 양식이 아직 설정되어 있지 않습니다.")
    cfg = _CFG[brand]

    rows = load_orders(start_date, end_date, brand=brand, archive_dir=archive_dir)
    rows.sort(key=lambda r: (r.get("order_date") or r.get("_run_date") or "", r.get("order_id") or ""))

    catalog_lookup = _load_catalog_lookup(cfg["catalog_file"]) if cfg.get("catalog_file") else {}

    wb = openpyxl.load_workbook(BASE_DIR / cfg["template_file"])
    ws = wb[cfg["template_sheet"]]

    year = _parse_date(start_date).year
    ws[cfg["title_cell"]] = cfg["title_template"].format(year=year)

    static_by_col = cfg.get("static_by_col", {})
    data_start = cfg["data_start_row"]
    letters = {k: openpyxl.utils.get_column_letter(v) for k, v in COL.items()}

    seen_orders = set()
    for i, rec in enumerate(rows):
        r = data_start + i
        order_id = rec.get("order_id")
        is_first = order_id not in seen_orders if order_id else True
        if order_id:
            seen_orders.add(order_id)

        od = _parse_date(rec.get("order_date")) or _parse_date(rec.get("_run_date"))

        ws.cell(row=r, column=COL["no"], value=i + 1)
        if od:
            ws.cell(row=r, column=COL["year"], value=od.year)
            ws.cell(row=r, column=COL["recv_month"], value=od.month)
            ws.cell(row=r, column=COL["recv_day"], value=od.day)
        ws.cell(row=r, column=COL["sales_channel"], value=rec.get("shop_name"))

        catalog_entry = catalog_lookup.get((rec.get("product_name"), rec.get("option") or ""))

        vendor = rec.get("vendor")
        if vendor:
            # 배송처(ship_office)는 실제로 물건을 포장·발송하는 협력사(예:
            # 투데이넛)를 그대로 쓰지만, 지급매입처(pay_vendor)는 그 협력사가
            # 아니라 실제로 대금을 지급하는 매입처가 따로 있는 경우가 있다
            # (예: 투데이넛 상품은 카탈로그상 매입처가 "해맑음푸드"). 카탈로그에
            # 그 상품의 pay_vendor가 등록돼 있으면 그걸 쓰고, 없으면 배송처와
            # 같다고 보고 vendor로 채운다.
            pay_vendor = (catalog_entry.get("pay_vendor") if catalog_entry else None) or vendor
            ws.cell(row=r, column=COL["pay_vendor"], value=pay_vendor)
            ws.cell(row=r, column=COL["ship_office"], value=vendor)

        ws.cell(row=r, column=COL["order_id"], value=order_id)
        ws.cell(row=r, column=COL["receiver_name"], value=rec.get("receiver_name"))
        ws.cell(row=r, column=COL["receiver_phone"], value=rec.get("receiver_phone"))
        ws.cell(row=r, column=COL["address"], value=rec.get("address"))
        ws.cell(row=r, column=COL["zipcode"], value=rec.get("zipcode"))
        ws.cell(row=r, column=COL["product_name"], value=rec.get("product_name"))
        ws.cell(row=r, column=COL["option"], value=rec.get("option"))

        if catalog_entry:
            ws.cell(row=r, column=COL["managed_name"], value=catalog_entry.get("managed_name"))
            unit_cost = catalog_entry.get("unit_cost")
            if unit_cost is not None:
                ws.cell(row=r, column=COL["unit_cost"], value=unit_cost)
                ws.cell(row=r, column=COL["cost_amount"], value=f"={letters['quantity']}{r}*{letters['unit_cost']}{r}")
                ws.cell(row=r, column=COL["profit"], value=f"={letters['settlement']}{r}-{letters['cost_amount']}{r}")
                ws.cell(
                    row=r, column=COL["margin"],
                    value=f'=IFERROR(({letters["settlement"]}{r}-{letters["cost_amount"]}{r})/{letters["settlement"]}{r}*100,0)',
                )

        ws.cell(row=r, column=COL["shipment_count"], value=1 if is_first else None)
        ws.cell(row=r, column=COL["quantity"], value=rec.get("quantity"))
        revenue = rec.get("payment_amount")
        ws.cell(row=r, column=COL["revenue"], value=revenue)
        if revenue is not None:
            ws.cell(row=r, column=COL["settlement"], value=f"={letters['revenue']}{r}*0.85")
        ws.cell(row=r, column=COL["delivery_message"], value=rec.get("delivery_message"))

        for col_letter, value in static_by_col.items():
            col_idx = openpyxl.utils.column_index_from_string(col_letter)
            ws.cell(row=r, column=col_idx, value=value)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    if not protect:
        wb.save(out_path)
        return out_path, len(rows)

    pwd = password or EXPORT_PASSWORD
    if not pwd:
        raise ValueError("protect=True인데 사용할 비밀번호가 없습니다 (config의 export_password 확인).")
    with tempfile.TemporaryDirectory(prefix="settlement-unprotected-") as tmp_dir:
        tmp_path = Path(tmp_dir) / out_path.name
        wb.save(tmp_path)
        protect_xlsx(tmp_path, out_path, pwd)
    return out_path, len(rows)
